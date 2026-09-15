"""Search Console "Performance on Search" 엑셀 → data/gsc_demand.json.

왜 필요한가 (2026-09-16 실측):
  - 2026-08-13 에 "Maps 사본" 등급(E·F) supplier 4,100 개를 noindex 로 돌렸다.
    그중 218 개가 실제로 구글 1~7위에 노출되던 페이지였다 — 노출 905 회,
    전체 클릭의 17%. 색인 페이지는 1,647 → 1,148 로 떨어졌다.
  - dead-lead 필터(전화·웹사이트 없음)는 SCG 시멘트 공장·NMB-Minebea 처럼
    구글이 클릭까지 주던 공급사를 지웠다.

두 규칙 모두 "이 페이지에 검색 가치가 있는가" 를 추측으로 판정했다. 추측 대신
구글이 실제로 노출해준 기록을 쓴다. 이 파일에 있는 경로는 등급·신호와 무관하게
색인·사이트맵·dead-lead 생존이 보장된다.

사용:
  python web-factory/scripts/import_gsc_pages.py <Performance-on-Search.xlsx> [...]

기존 기록과 합친다 (경로별 최대 노출·클릭, 최고 순위) — 3개월 창이 밀려나도
한 번 증명된 수요는 남는다.
"""
from __future__ import annotations

import json
import sys
from datetime import date
from pathlib import Path
from urllib.parse import unquote, urlparse

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "data" / "gsc_demand.json"


def read_pages(xlsx: Path) -> dict[str, dict]:
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    if "Pages" not in wb.sheetnames:
        sys.exit(f"{xlsx.name}: 'Pages' 시트가 없습니다 — Performance 내보내기 파일인지 확인")
    rows = list(wb["Pages"].iter_rows(values_only=True))[1:]
    out: dict[str, dict] = {}
    for url, clicks, impr, _ctr, pos in rows:
        if not url:
            continue
        path = unquote(urlparse(url).path).rstrip("/") or "/"
        out[path] = {"impressions": int(impr or 0), "clicks": int(clicks or 0),
                     "best_position": round(float(pos or 0), 2)}
    return out


def main(argv: list[str]) -> int:
    if not argv:
        print(__doc__)
        return 1
    db = {"updated": None, "pages": {}}
    if OUT.exists():
        db = json.loads(OUT.read_text(encoding="utf-8"))
    pages: dict[str, dict] = db.get("pages", {})
    before = len(pages)
    for arg in argv:
        for path, rec in read_pages(Path(arg)).items():
            cur = pages.get(path)
            if not cur:
                pages[path] = rec
                continue
            cur["impressions"] = max(cur["impressions"], rec["impressions"])
            cur["clicks"] = max(cur["clicks"], rec["clicks"])
            if rec["best_position"] and (not cur["best_position"]
                                         or rec["best_position"] < cur["best_position"]):
                cur["best_position"] = rec["best_position"]
    db = {"updated": date.today().isoformat(),
          "note": "Search Console 에서 노출이 기록된 경로. scripts/import_gsc_pages.py 가 갱신한다.",
          "pages": dict(sorted(pages.items()))}
    OUT.write_text(json.dumps(db, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    sup = sum(1 for p in pages if p.startswith("/supplier/"))
    print(f"gsc_demand.json: {before} → {len(pages)} paths ({sup} supplier)")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
